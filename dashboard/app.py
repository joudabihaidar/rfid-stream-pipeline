import sys
import logging
import time
from pathlib import Path

import pandas as pd
import plotly.express as px
import streamlit as st

# Add project root AND src/ to sys.path so all internal imports resolve
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from src.database import (
    get_all_sessions,
    get_raw_rows,
    get_events_for_session,
    get_all_anomalies,
    get_dwell_times,
    get_door_usage,
    get_zone_transition_paths,
    get_anomaly_summary,
)
from src.pipeline import stream_pipeline_events_from_rows

EXCEL_PATH = str(PROJECT_ROOT / "data" / "raw_data.xlsx")
LOG_PATH   = PROJECT_ROOT / "data" / "pipeline.log"

# ═══════════════════════════════════════════════════════════════
# PAGE CONFIG
# ═══════════════════════════════════════════════════════════════

st.set_page_config(
    page_title="RFID Building Access Tracker",
    page_icon="🏢",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ═══════════════════════════════════════════════════════════════
# STREAMLIT LOG HANDLER
# Writes pipeline log lines into a Streamlit placeholder in real time
# ═══════════════════════════════════════════════════════════════

class StreamlitLogHandler(logging.Handler):
    """
    Custom logging handler that appends log lines to a Streamlit
    st.empty() placeholder as the pipeline runs.
    Only shows the last MAX_LINES lines so the box doesn't grow forever.
    """
    MAX_LINES = 60

    def __init__(self, placeholder):
        super().__init__()
        self.placeholder = placeholder
        self.lines       = []

    def emit(self, record):
        self.lines.append(self.format(record))
        self.placeholder.code(
            "\n".join(self.lines[-self.MAX_LINES:]),
            language="log",
        )


# ═══════════════════════════════════════════════════════════════
# DATA LOADING (cached — refreshes every 60 seconds)
# ═══════════════════════════════════════════════════════════════

@st.cache_data(ttl=60)
def load_sessions():
    return pd.DataFrame(get_all_sessions())

@st.cache_data(ttl=60)
def load_anomalies():
    return pd.DataFrame(get_all_anomalies())

@st.cache_data(ttl=60)
def load_dwell_times():
    return pd.DataFrame(get_dwell_times())

@st.cache_data(ttl=60)
def load_door_usage():
    return pd.DataFrame(get_door_usage())

@st.cache_data(ttl=60)
def load_zone_paths():
    return pd.DataFrame(get_zone_transition_paths())

@st.cache_data(ttl=60)
def load_anomaly_summary():
    return pd.DataFrame(get_anomaly_summary())

sessions_df  = load_sessions()
anomalies_df = load_anomalies()
dwell_df     = load_dwell_times()
door_df      = load_door_usage()
zone_df      = load_zone_paths()
anomaly_sum  = load_anomaly_summary()

# ═══════════════════════════════════════════════════════════════
# SIDEBAR
# ═══════════════════════════════════════════════════════════════

with st.sidebar:
    st.title("🏢 RFID Tracker")
    st.caption("Real-time building access monitoring")
    st.divider()

    session_ids = sessions_df["session_id"].tolist() if not sessions_df.empty else []

    selected_session = st.selectbox(
        "Session",
        options=session_ids,
        index=0 if session_ids else None,
        help="Select a test case session to inspect",
    )

    st.divider()
    st.subheader("Live replay")
    st.caption(
        "Feeds raw RFID rows into the detection pipeline one at a time. "
        "Events appear the moment the algorithm confirms them."
    )
    replay_speed = st.slider(
        "Delay between raw rows (s)",
        min_value=0.01,
        max_value=0.5,
        value=0.05,
        step=0.01,
        help="Lower = faster. At 0.05s a 100-row session takes ~5 seconds.",
    )
    start_replay = st.button(
        "▶ Start Replay",
        use_container_width=True,
        type="primary",
    )

    st.divider()
    total_events    = int(sessions_df["total_entries"].sum() + sessions_df["total_exits"].sum()) if not sessions_df.empty else 0
    total_anomalies = len(anomalies_df) if not anomalies_df.empty else 0
    st.caption(
        f"{len(sessions_df)} sessions · "
        f"{total_events} events · "
        f"{total_anomalies} anomalies"
    )


# ═══════════════════════════════════════════════════════════════
# SHARED RENDER HELPERS
# ═══════════════════════════════════════════════════════════════

def render_metrics(displayed_events, metrics_area, n_anomalies):
    entries = sum(1 for e in displayed_events if e.get("event_type") == "ENTRY")
    exits   = sum(1 for e in displayed_events if e.get("event_type") == "EXIT")
    inside  = max(0, entries - exits)

    with metrics_area.container():
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Currently inside", inside)
        c2.metric("Entries",          entries)
        c3.metric("Exits",            exits)
        c4.metric(
            "Anomalies flagged", n_anomalies,
            delta="review required" if n_anomalies > 0 else None,
            delta_color="inverse",
        )


def render_feed(displayed_events, feed_area):
    df = pd.DataFrame(displayed_events)
    rename = {
        "event_type": "Event",
        "t0":         "T0 (s)",
        "door":       "Door",
        "from_zone":  "From zone",
        "to_zone":    "To zone",
        "peak_rssi":  "Peak RSSI",
        "dwell_time": "Dwell (s)",
    }
    df = df.rename(columns=rename)
    for col in ["From zone", "To zone"]:
        if col in df.columns:
            df[col] = df[col].str.split("__").str[-1]

    cols = ["Event", "T0 (s)", "Door", "From zone", "To zone", "Dwell (s)", "Peak RSSI"]
    cols = [c for c in cols if c in df.columns]

    with feed_area.container():
        st.dataframe(df[cols], width="stretch", height=320)


# ═══════════════════════════════════════════════════════════════
# TABS
# ═══════════════════════════════════════════════════════════════

tab1, tab2, tab3, tab4 = st.tabs([
    "📡 Live Feed",
    "📊 Access Patterns",
    "⚠️ Anomalies",
    "⚙️ Pipeline",
])


# ═══════════════════════════════════════════════════════════════
# TAB 1 — LIVE FEED
# ═══════════════════════════════════════════════════════════════

with tab1:
    if not selected_session:
        st.info("Select a session from the sidebar.")
    else:
        session_row = sessions_df[
            sessions_df["session_id"] == selected_session
        ].iloc[0]

        st.subheader(f"Session: {selected_session}")
        st.caption(
            f"Date: {session_row.get('session_date', 'N/A')}  ·  "
            f"EPC: ...{str(session_row.get('epc', ''))[-8:]}  ·  "
            f"Entry signal: {session_row.get('entry_rssi_strength', 'N/A')}  ·  "
            f"Ghost ratio: {float(session_row.get('ghost_read_ratio') or 0):.1%}"
        )
        st.divider()

        metrics_area = st.empty()
        st.markdown("**Event log**")
        feed_area = st.empty()

        # Pre-load anomaly info for this session
        session_anomalies = (
            anomalies_df[anomalies_df["session_id"] == selected_session]
            if not anomalies_df.empty else pd.DataFrame()
        )
        n_anomalies = len(session_anomalies)
        anomaly_t0s = set(
            session_anomalies["t0"].dropna().astype(int).tolist()
        ) if not session_anomalies.empty else set()

        if start_replay:
            # ── Real-time replay ──────────────────────────────────────
            # Feeds raw rows from DB one by one into the pipeline.
            # Events are yielded the moment the algorithm confirms them.
            db_rows   = get_raw_rows(selected_session)
            displayed = []

            for event in stream_pipeline_events_from_rows(
                db_rows, delay=replay_speed
            ):
                displayed.append(event)
                render_metrics(displayed, metrics_area, n_anomalies)
                render_feed(displayed, feed_area)

                event_type = event.get("event_type", "")
                t0_val     = event.get("t0")

                if event_type == "SESSION_ENDED_INSIDE":
                    st.toast("⚠️ Session ended inside — no exit detected", icon="🚨")
                elif t0_val is not None and int(t0_val) in anomaly_t0s:
                    st.toast(f"⚠️ {event_type} detected at T0={t0_val}", icon="🚨")

        else:
            # ── Static view before replay ─────────────────────────────
            session_events = get_events_for_session(selected_session)
            if session_events:
                normalized = [
                    {**e, "event_type": e.get("event_type", e.get("event", ""))}
                    for e in session_events
                ]
                render_metrics(normalized, metrics_area, n_anomalies)
                render_feed(normalized, feed_area)
            else:
                st.info("No events found. Run the pipeline first.")


# ═══════════════════════════════════════════════════════════════
# TAB 2 — ACCESS PATTERNS
# ═══════════════════════════════════════════════════════════════

with tab2:
    if sessions_df.empty:
        st.info("No data available. Run the pipeline from the ⚙️ Pipeline tab.")
    else:
        total_entries     = int(sessions_df["total_entries"].sum())
        total_exits       = int(sessions_df["total_exits"].sum())
        total_transitions = int(sessions_df["total_transitions"].sum())
        avg_dwell         = dwell_df["dwell_time"].mean() if not dwell_df.empty else 0
        max_dwell         = dwell_df["dwell_time"].max()  if not dwell_df.empty else 0

        c1, c2, c3, c4, c5 = st.columns(5)
        c1.metric("Total entries",    total_entries)
        c2.metric("Total exits",      total_exits)
        c3.metric("Zone transitions", total_transitions)
        c4.metric("Avg dwell time",   f"{avg_dwell:.0f}s")
        c5.metric("Max dwell time",   f"{max_dwell:.0f}s")

        st.divider()

        col1, col2 = st.columns(2)

        with col1:
            st.subheader("Entries and exits per session")
            chart_df = sessions_df[
                ["session_id", "total_entries", "total_exits"]
            ].copy()
            chart_df["session_id"] = chart_df["session_id"].str[-14:]
            melted = chart_df.melt(
                id_vars="session_id",
                value_vars=["total_entries", "total_exits"],
                var_name="type", value_name="count",
            )
            melted["type"] = melted["type"].map({
                "total_entries": "Entries",
                "total_exits":   "Exits",
            })
            fig = px.bar(
                melted, x="session_id", y="count",
                color="type", barmode="group",
                labels={"session_id": "Session", "count": "Count", "type": ""},
                color_discrete_map={"Entries": "#1d9e75", "Exits": "#D85A30"},
            )
            fig.update_layout(
                margin=dict(t=10, b=10), height=300,
                legend=dict(orientation="h"),
            )
            st.plotly_chart(fig, use_container_width=True)

        with col2:
            st.subheader("Door usage — entries vs exits")
            if not door_df.empty:
                fig = px.bar(
                    door_df, x="door", y="count",
                    color="event_type", barmode="group",
                    labels={"door": "Door", "count": "Count", "event_type": ""},
                    color_discrete_map={"ENTRY": "#1d9e75", "EXIT": "#D85A30"},
                )
                fig.update_layout(
                    margin=dict(t=10, b=10), height=300,
                    legend=dict(orientation="h"),
                )
                st.plotly_chart(fig, use_container_width=True)
            else:
                st.info("No door usage data.")

        st.divider()

        col3, col4 = st.columns(2)

        with col3:
            st.subheader("Dwell time distribution")
            if not dwell_df.empty:
                fig = px.histogram(
                    dwell_df, x="dwell_time", nbins=8,
                    labels={
                        "dwell_time": "Dwell time (seconds)",
                        "count":      "Sessions",
                    },
                    color_discrete_sequence=["#534AB7"],
                )
                fig.update_layout(margin=dict(t=10, b=10), height=300)
                st.plotly_chart(fig, use_container_width=True)
            else:
                st.info("No dwell time data available.")

        with col4:
            st.subheader("Zone transition paths")
            if not zone_df.empty:
                display_zone = zone_df.copy()
                display_zone["from_zone"] = (
                    display_zone["from_zone"].str.split("__").str[-1]
                )
                display_zone["to_zone"] = (
                    display_zone["to_zone"].str.split("__").str[-1]
                )
                display_zone.columns = ["From", "To", "Count"]
                st.dataframe(display_zone, width="stretch", height=300)
            else:
                st.info("No zone transitions recorded.")

        st.divider()

        st.subheader("Data quality — ghost read ratio per session")
        st.caption(
            "Proportion of reads that were Out reads while the person "
            "was confirmed inside. Higher = noisier signal."
        )
        if "ghost_read_ratio" in sessions_df.columns:
            ghost_df = sessions_df[[
                "session_id", "ghost_read_ratio", "entry_rssi_strength"
            ]].copy()
            ghost_df["session_id"] = ghost_df["session_id"].str[-14:]
            ghost_df["ghost_read_ratio"] = ghost_df["ghost_read_ratio"].fillna(0)
            ghost_df["entry_rssi_strength"] = (
                ghost_df["entry_rssi_strength"].fillna("unknown")
            )
            fig = px.bar(
                ghost_df,
                x="session_id", y="ghost_read_ratio",
                color="entry_rssi_strength",
                labels={
                    "session_id":          "Session",
                    "ghost_read_ratio":    "Ghost read ratio",
                    "entry_rssi_strength": "Entry signal",
                },
                color_discrete_map={
                    "strong":   "#1d9e75",
                    "moderate": "#BA7517",
                    "weak":     "#D85A30",
                    "unknown":  "#888780",
                },
            )
            fig.update_layout(
                margin=dict(t=10, b=10), height=280,
                yaxis_tickformat=".0%",
                legend=dict(orientation="h"),
            )
            st.plotly_chart(fig, use_container_width=True)


# ═══════════════════════════════════════════════════════════════
# TAB 3 — ANOMALIES
# ═══════════════════════════════════════════════════════════════

with tab3:
    if anomalies_df.empty:
        st.success("✅ No anomalies detected across all sessions.")
    else:
        total_anom       = len(anomalies_df)
        flagged_sessions = anomalies_df["session_id"].nunique()
        clean_sessions   = len(sessions_df) - flagged_sessions

        c1, c2, c3 = st.columns(3)
        c1.metric("Total anomalies",         total_anom)
        c2.metric("Sessions with anomalies", flagged_sessions)
        c3.metric("Clean sessions",          clean_sessions)

        st.divider()

        col1, col2 = st.columns(2)

        with col1:
            st.subheader("Anomalies by type")
            if not anomaly_sum.empty:
                fig = px.bar(
                    anomaly_sum, x="anomaly_type", y="count",
                    labels={
                        "anomaly_type": "Anomaly type",
                        "count":        "Count",
                    },
                    color_discrete_sequence=["#D85A30"],
                )
                fig.update_layout(margin=dict(t=10, b=10), height=300)
                st.plotly_chart(fig, use_container_width=True)

        with col2:
            st.subheader("Anomalies per session")
            per_session = (
                anomalies_df.groupby("session_id")["anomaly_type"]
                .count()
                .reset_index()
                .rename(columns={"anomaly_type": "count"})
            )
            per_session["session_id"] = per_session["session_id"].str[-14:]
            fig = px.bar(
                per_session, x="session_id", y="count",
                labels={"session_id": "Session", "count": "Anomaly count"},
                color_discrete_sequence=["#993C1D"],
            )
            fig.update_layout(margin=dict(t=10, b=10), height=300)
            st.plotly_chart(fig, use_container_width=True)

        st.divider()

        st.subheader("Session overview")
        overview = sessions_df[[
            "session_id", "session_date", "total_entries", "total_exits",
            "total_transitions", "ghost_read_ratio",
            "entry_rssi_strength", "has_anomaly",
        ]].copy()
        overview["has_anomaly"] = overview["has_anomaly"].map(
            {1: "⚠️ Yes", 0: "✅ No"}
        )
        overview["ghost_read_ratio"] = overview["ghost_read_ratio"].apply(
            lambda x: f"{float(x):.1%}" if x is not None else "N/A"
        )
        overview.columns = [
            "Session", "Date", "Entries", "Exits",
            "Transitions", "Ghost ratio", "Entry signal", "Anomaly",
        ]
        st.dataframe(overview, width="stretch", height=280)

        st.divider()

        st.subheader("Full anomaly log")
        log = anomalies_df[[
            "session_id", "epc", "anomaly_type", "t0", "value", "note"
        ]].copy()
        log["epc"]        = log["epc"].str[-8:]
        log["session_id"] = log["session_id"].str[-14:]
        log.columns       = ["Session", "EPC", "Type", "T0 (s)", "Value", "Note"]
        st.dataframe(log, width="stretch", height=400)


# ═══════════════════════════════════════════════════════════════
# TAB 4 — PIPELINE
# ═══════════════════════════════════════════════════════════════

with tab4:
    st.subheader("Run pipeline")
    st.caption(
        "For each session (Excel sheet): "
        "**Stage 1** extracts raw RFID reads into the database · "
        "**Stage 2** runs the detection algorithm · "
        "**Stage 3** detects anomalies and computes data quality · "
        "**Stage 4** stores all results. "
        "Metrics update after each session. Logs stream in real time."
    )

    run_btn = st.button(
        "▶ Run full pipeline",
        type="primary",
        use_container_width=True,
    )

    if run_btn:
        import openpyxl
        from src.main     import process_sheet
        from src.database import create_schema

        # ── Live metric placeholders ──────────────────────────────
        st.markdown("**Live progress**")
        m1, m2, m3, m4 = st.columns(4)
        entries_box    = m1.empty()
        exits_box      = m2.empty()
        anomalies_box  = m3.empty()
        sessions_box   = m4.empty()

        entries_box.metric("Entries detected",    0)
        exits_box.metric("Exits detected",        0)
        anomalies_box.metric("Anomalies found",   0)
        sessions_box.metric("Sessions processed", "0 / ?")

        # ── Live log placeholder ──────────────────────────────────
        st.markdown("**Live logs**")
        log_placeholder = st.empty()

        # ── Attach logging handlers ───────────────────────────────
        # StreamlitLogHandler → updates the placeholder in real time
        # FileHandler         → writes to pipeline.log
        LOG_PATH.parent.mkdir(parents=True, exist_ok=True)

        streamlit_handler = StreamlitLogHandler(log_placeholder)
        file_handler      = logging.FileHandler(LOG_PATH, encoding="utf-8")

        formatter = logging.Formatter(
            "%(asctime)s  %(levelname)-8s  %(message)s",
            datefmt="%H:%M:%S",
        )
        streamlit_handler.setFormatter(formatter)
        file_handler.setFormatter(formatter)

        root_logger = logging.getLogger()
        root_logger.setLevel(logging.DEBUG)
        root_logger.addHandler(streamlit_handler)
        root_logger.addHandler(file_handler)
        logging.getLogger("openpyxl").setLevel(logging.WARNING)

        # ── Run pipeline session by session via process_sheet ─────
        # process_sheet() is the exact same function main.py uses —
        # one source of truth, identical logic, identical logs.
        create_schema()

        try:
            wb          = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
            sheet_names = wb.sheetnames
            wb.close()
        except Exception as e:
            st.error(f"Could not open Excel file: {e}")
            sheet_names = []

        total_entries   = 0
        total_exits     = 0
        total_anomalies = 0

        for i, sheet_name in enumerate(sheet_names, start=1):
            sessions_box.metric(
                "Sessions processed",
                f"{i - 1} / {len(sheet_names)}",
            )

            # Single call — same function, same logic, same logs as main.py
            events, anomalies = process_sheet(str(EXCEL_PATH), sheet_name)

            total_entries   += sum(1 for e in events if e.get("event") == "ENTRY")
            total_exits     += sum(1 for e in events if e.get("event") == "EXIT")
            total_anomalies += len(anomalies)

            entries_box.metric("Entries detected",    total_entries)
            exits_box.metric("Exits detected",        total_exits)
            anomalies_box.metric("Anomalies found",   total_anomalies)
            sessions_box.metric(
                "Sessions processed",
                f"{i} / {len(sheet_names)}",
            )

        # Detach handlers
        root_logger.removeHandler(streamlit_handler)
        root_logger.removeHandler(file_handler)
        file_handler.close()

        st.success(
            f"✅ Pipeline complete — "
            f"{len(sheet_names)} sessions · "
            f"{total_entries} entries · "
            f"{total_exits} exits · "
            f"{total_anomalies} anomalies"
        )

        st.cache_data.clear()
        time.sleep(1)
        st.rerun()

    # ── Log file viewer ───────────────────────────────────────────
    st.divider()
    st.subheader("Pipeline log file")
    st.caption(f"Path: data/pipeline.log")

    if LOG_PATH.exists():
        log_text = LOG_PATH.read_text(encoding="utf-8")
        lines    = log_text.strip().split("\n")

        col1, col2 = st.columns([3, 1])
        with col2:
            n_lines = st.selectbox(
                "Show last N lines",
                options=[50, 100, 200, 500],
                index=0,
            )
        with col1:
            st.caption(f"{len(lines)} total lines in log file")

        st.code("\n".join(lines[-n_lines:]), language="log")
    else:
        st.info("No log file yet. Run the pipeline first.")